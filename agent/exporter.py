"""Markdown -> DOCX converter for board pre-reads.

Self-contained lightweight markdown parser plus python-docx renderer.
Tuned for the Brock pre-read format:
  # / ## / ### headings, **bold**, *italic*, > blockquotes (with flag markers),
  | pipe | tables |, numbered/bulleted lists.

The output is a Word document a trustee can carry into a meeting:
serif title, navy headings, real Word tables, color-coded flag callouts.
"""

from __future__ import annotations

import io
import re
from typing import List, Tuple

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches, Cm

# --------------------------------------------------------------------------- #
# Color palette (matches the web blockquote palette)
# --------------------------------------------------------------------------- #
NAVY = RGBColor(0x1A, 0x3A, 0x6B)
INK = RGBColor(0x14, 0x14, 0x13)
INK_SOFT = RGBColor(0x2B, 0x2A, 0x27)
MUTE = RGBColor(0x6B, 0x68, 0x60)
RED = RGBColor(0xC4, 0x5A, 0x4A)
ORANGE = RGBColor(0xD9, 0x77, 0x57)
GREEN = RGBColor(0x78, 0x8C, 0x5D)
HEADER_BG = "E8E6DC"  # cream-2 hex for table header shading
RED_BG = "FCE7E2"
ORANGE_BG = "FCEBE0"
GREEN_BG = "ECF0E0"

BODY_FONT = "Cambria"


# --------------------------------------------------------------------------- #
# Low-level XML helpers (paragraph shading + borders)
# --------------------------------------------------------------------------- #
def _set_cell_shading(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _set_paragraph_shading(paragraph, hex_color: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    p_pr.append(shd)


def _set_left_border(paragraph, hex_color: str, width_pt: int = 18) -> None:
    """Add a colored left border to a paragraph (used for flag callouts)."""
    p_pr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(width_pt))  # eighths of a point
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), hex_color)
    pbdr.append(left)
    p_pr.append(pbdr)


# --------------------------------------------------------------------------- #
# Inline parsing: bold / italic / code / "strikethrough-with-suffix"
# --------------------------------------------------------------------------- #
_INLINE_RE = re.compile(
    r"(\*\*[^*\n]+\*\*|\*[^*\n]+\*|`[^`\n]+`|~~[^~\n]+~~)"
)


def _add_inline_runs(paragraph, text: str, *, base_size: float = 11.0,
                     base_color: RGBColor | None = None,
                     italic: bool = False, bold: bool = False) -> None:
    """Split *text* on bold/italic/code markers and add styled runs."""
    text = text.replace("\\*", "\u0001").replace("\\_", "\u0002")
    parts = _INLINE_RE.split(text)
    for part in parts:
        if not part:
            continue
        b, i, code, strike = bold, italic, False, False
        content = part
        if part.startswith("**") and part.endswith("**"):
            b = True
            content = part[2:-2]
        elif part.startswith("*") and part.endswith("*"):
            i = True
            content = part[1:-1]
        elif part.startswith("`") and part.endswith("`"):
            code = True
            content = part[1:-1]
        elif part.startswith("~~") and part.endswith("~~"):
            strike = True
            content = part[2:-2]
        content = content.replace("\u0001", "*").replace("\u0002", "_")
        run = paragraph.add_run(content)
        run.bold = b
        run.italic = i
        run.font.size = Pt(base_size)
        run.font.name = "Consolas" if code else BODY_FONT
        if strike:
            run.font.strike = True
            run.font.color.rgb = MUTE
        elif base_color is not None:
            run.font.color.rgb = base_color
        else:
            run.font.color.rgb = INK_SOFT


# --------------------------------------------------------------------------- #
# Block-level parsing (line-by-line state machine)
# --------------------------------------------------------------------------- #
def _parse_table_row(line: str) -> List[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def _is_table_separator(line: str) -> bool:
    s = line.strip()
    if not s.startswith("|"):
        return False
    cells = _parse_table_row(s)
    if not cells:
        return False
    return all(re.fullmatch(r":?-+:?", c.replace(" ", "")) for c in cells if c)


def _flag_kind(text: str) -> str | None:
    """Detect RED FLAG / WATCH / POSITIVE in a blockquote line."""
    t = text.upper()
    if "RED FLAG" in t:
        return "red"
    if "WATCH" in t and "WATCH" in text.upper():
        # Only treat as flag if marker word is bold-formatted or at start
        if re.search(r"\*\*WATCH\b", text) or t.lstrip().startswith("WATCH"):
            return "orange"
    if "POSITIVE" in t:
        if re.search(r"\*\*POSITIVE\b", text) or t.lstrip().startswith("POSITIVE"):
            return "green"
    return None


# --------------------------------------------------------------------------- #
# Renderers for individual block types
# --------------------------------------------------------------------------- #
def _add_heading(doc: Document, text: str, level: int) -> None:
    sizes = {1: 26, 2: 18, 3: 14, 4: 12}
    size = sizes.get(level, 11)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    # Strip surrounding ** that some templates wrap headings in
    text = text.strip()
    if text.startswith("**") and text.endswith("**"):
        text = text[2:-2]
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.name = BODY_FONT
    run.font.color.rgb = NAVY


def _add_paragraph(doc: Document, text: str) -> None:
    if not text.strip():
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    _add_inline_runs(p, text)


def _add_list_item(doc: Document, text: str, ordered: bool, number: int = 1,
                   italic: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.space_after = Pt(2)
    bullet = f"{number}." if ordered else "\u2022"
    head = p.add_run(f"{bullet}  ")
    head.font.size = Pt(11)
    head.font.name = BODY_FONT
    head.font.color.rgb = INK_SOFT
    _add_inline_runs(p, text, italic=italic)


def _add_blockquote(doc: Document, lines: List[str]) -> None:
    text = "\n".join(lines).strip()
    kind = _flag_kind(text)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.15)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    if kind == "red":
        _set_paragraph_shading(p, RED_BG)
        _set_left_border(p, "C45A4A")
    elif kind == "orange":
        _set_paragraph_shading(p, ORANGE_BG)
        _set_left_border(p, "D97757")
    elif kind == "green":
        _set_paragraph_shading(p, GREEN_BG)
        _set_left_border(p, "788C5D")
    else:
        _set_left_border(p, "B0AEA5")
    _add_inline_runs(p, text, base_color=INK_SOFT)


def _add_table(doc: Document, header: List[str], rows: List[List[str]]) -> None:
    cols = max(len(header), max((len(r) for r in rows), default=0))
    if cols == 0:
        return
    table = doc.add_table(rows=1 + len(rows), cols=cols)
    table.style = "Light Grid Accent 1"
    # Header
    hdr_cells = table.rows[0].cells
    for i in range(cols):
        cell = hdr_cells[i]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_shading(cell, HEADER_BG)
        cell.text = ""
        p = cell.paragraphs[0]
        text = header[i] if i < len(header) else ""
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.name = BODY_FONT
        run.font.color.rgb = NAVY
    # Body
    for r_idx, row in enumerate(rows, start=1):
        body_cells = table.rows[r_idx].cells
        for c_idx in range(cols):
            cell = body_cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            text = row[c_idx] if c_idx < len(row) else ""
            _add_inline_runs(p, text, base_size=10.0)
    doc.add_paragraph()  # spacer


# --------------------------------------------------------------------------- #
# Main entry point
# --------------------------------------------------------------------------- #
def markdown_to_docx(md_text: str, title: str = "Board Meeting Pre-Read") -> bytes:
    doc = Document()

    # Default style: Cambria 11pt
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(11)

    # Page margins (slightly tighter than default for a denser pre-read)
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    # ---- Title ----
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(18)
    t_run = title_p.add_run(title)
    t_run.bold = True
    t_run.font.name = BODY_FONT
    t_run.font.size = Pt(22)
    t_run.font.color.rgb = NAVY

    # ---- Footer ----
    footer = doc.sections[0].footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f_run = footer_p.add_run(
        "Prepared by Texas Governance Agent  \u00b7  Claude Opus 4.7"
    )
    f_run.font.name = BODY_FONT
    f_run.font.size = Pt(9)
    f_run.italic = True
    f_run.font.color.rgb = MUTE

    # ---- Body parse ----
    lines = md_text.splitlines()
    i = 0
    ol_counter = 0  # running counter for numbered lists
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Blank line
        if not stripped:
            ol_counter = 0
            i += 1
            continue

        # Horizontal rule
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            doc.add_paragraph()  # cheap visual break
            i += 1
            continue

        # Heading
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            _add_heading(doc, m.group(2), level)
            ol_counter = 0
            i += 1
            continue

        # Blockquote (consume contiguous > lines)
        if stripped.startswith(">"):
            block: List[str] = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                block.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            _add_blockquote(doc, block)
            ol_counter = 0
            continue

        # Table (header row, separator row, body rows)
        if stripped.startswith("|") and (i + 1) < len(lines) and \
                _is_table_separator(lines[i + 1]):
            header = _parse_table_row(lines[i])
            i += 2
            rows: List[List[str]] = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(_parse_table_row(lines[i]))
                i += 1
            _add_table(doc, header, rows)
            ol_counter = 0
            continue

        # Ordered list item
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            ol_counter += 1
            text = m.group(2)
            italic = text.startswith("*") and text.endswith("*") and \
                not text.startswith("**")
            _add_list_item(doc, text, ordered=True, number=ol_counter,
                           italic=italic)
            i += 1
            continue

        # Unordered list item
        m = re.match(r"^[-*+]\s+(.*)$", stripped)
        if m:
            _add_list_item(doc, m.group(1), ordered=False)
            ol_counter = 0
            i += 1
            continue

        # Paragraph (gather contiguous non-special lines)
        para_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt:
                break
            if re.match(r"^(#{1,6}\s|>\s|>$|\||[-*+]\s|\d+\.\s)", nxt):
                break
            if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", nxt):
                break
            para_lines.append(nxt)
            i += 1
        _add_paragraph(doc, " ".join(para_lines))
        ol_counter = 0

    # ---- Serialize ----
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
